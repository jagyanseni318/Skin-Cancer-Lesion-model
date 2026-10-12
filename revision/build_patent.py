"""Builds the patent specification (DOCX) from the executed experimental results."""
import json
import numpy as np, pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from common import *
import ms_results_xai as RX
from ms_results_cls import pm

RES = OUT / 'results'; PD = OUT / 'patent'; XA = RX.XA; PA = RX.PA; K = len(XA); ac = RX.ac
S = json.load(open(RES / 'cls_summary.json')); agg = S['agg']; R = XA[42]
gap = ac('perm_null', 'rho_obs') - ac('perm_null', 'rho_null_mean')
rho_obs = ac('perm_null', 'rho_obs'); rho_null = ac('perm_null', 'rho_null_mean'); e_med = ac('enrichment_vs_null', 'all', 'median_e_shuf'); e_area = ac('size_measures', 'e_shuf', 'spearman_area', 'rho')
ig = ac('faithfulness', 'mean', 'IG', 'median'); dg = ac('faithfulness', 'mean', 'DG', 'median'); ig_f = 100 * ac('faithfulness', 'mean', 'frac_IG_pos')
e_ig = ac('faithfulness', 'mean', 'rho_eshuf_IG', 'rho'); l_ig = ac('faithfulness', 'mean', 'rho_lsas_IG', 'rho'); l_dg = ac('faithfulness', 'mean', 'rho_lsas_DG', 'rho')
sts_e = ac('sts', 'rho_eshuf', 'rho'); sts_l = ac('sts', 'rho_lsas', 'rho'); sts_m = ac('sts', 'median')
C5 = np.array([PA[s]['C_HAM_test_5class_images_5way_output']['acc'] for s in PA]); E5 = np.array([PA[s]['E_PAD_mapped_5way_output']['acc'] for s in PA]); Pv = np.array([PA[s]['prevalence_reweighted_HAM_acc_5way'] for s in PA])
A7 = np.array([PA[s]['A_HAM_test_7class']['acc'] for s in PA])
sm = R['size_measures']; sim = json.load(open(RES / 'lsas_scenarios_simulation.json'))
DGp = RX.DG; dgs = sorted(DGp)
dcauc = np.array([DGp[s]['domain_classifier_auroc']['mean'] for s in dgs])
lad = lambda k, m: np.array([DGp[s]['adaptation_ladder'][k][m][0] for s in dgs])

doc = Document(); sec = doc.sections[0]; sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
sec.left_margin = Cm(3.0); sec.right_margin = Cm(2.2); sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.2)
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12); st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(6)
para_no = [0]


def footer_page_numbers():
    p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for t, txt in [('begin', None), ('instr', 'PAGE'), ('end', None)]:
        r = p.add_run()
        if t == 'instr':
            el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = txt
        else:
            el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), t)
        r._r.append(el)


footer_page_numbers()


def H(text, size=12, center=False, before=12, page_break=False):
    if page_break: doc.add_page_break()
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.keep_with_next = True
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def P(text, numbered=True, indent=True):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if numbered:
        para_no[0] += 1; r = p.add_run(f'[{para_no[0]:04d}]  '); r.bold = True
    for part in text.split('**'):
        pass
    import re
    for i, part in enumerate(re.split(r'(\*\*[^*]+\*\*|__[^_]+__)', text)):
        if part.startswith('**'): rr = p.add_run(part[2:-2]); rr.bold = True
        elif part.startswith('__'): rr = p.add_run(part[2:-2]); rr.italic = True
        else: p.add_run(part)
    return p


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color); tcPr.append(shd)


def T(header, rows, caption):
    c = doc.add_paragraph(); r = c.add_run(caption); r.bold = True; r.font.size = Pt(10.5); c.paragraph_format.keep_with_next = True; c.paragraph_format.line_spacing = 1.0
    tb = doc.add_table(rows=1, cols=len(header)); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = tb.rows[0].cells[i]; cell.text = ''; rr = cell.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(9.5); shade(cell, 'E7E6E6'); cell.paragraphs[0].paragraph_format.line_spacing = 1.0
    for row in rows:
        cells = tb.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''; rr = cells[i].paragraphs[0].add_run(str(v)); rr.font.size = Pt(9.5); cells[i].paragraphs[0].paragraph_format.line_spacing = 1.0
    doc.add_paragraph().paragraph_format.line_spacing = 1.0


def FIG(path, caption, width=15.0):
    doc.add_page_break()
    doc.add_picture(str(path), width=Cm(width)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = c.add_run(caption); r.bold = True; r.font.size = Pt(11); c.paragraph_format.line_spacing = 1.0


# ============================================================================== TITLE PAGE
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = t.add_run('COMPLETE SPECIFICATION'); r.bold = True; r.font.size = Pt(13)
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = t.add_run('(DRAFT FOR REVIEW BY A REGISTERED PATENT AGENT / ATTORNEY)'); r.italic = True; r.font.size = Pt(10.5)
for _ in range(2): doc.add_paragraph()
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('SYSTEM AND METHOD FOR NULL-CONTROLLED RELIABILITY ASSESSMENT OF SALIENCY EXPLANATIONS IN DERMOSCOPIC SKIN-LESION CLASSIFICATION'); r.bold = True; r.font.size = Pt(17)
t.paragraph_format.line_spacing = 1.25
for _ in range(2): doc.add_paragraph()
tb = doc.add_table(rows=0, cols=2); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
def rowt(a, b):
    cells = tb.add_row().cells; cells[0].text = ''; cells[1].text = ''
    ra = cells[0].paragraphs[0].add_run(a); ra.bold = True; ra.font.size = Pt(10.5); rb = cells[1].paragraphs[0].add_run(b); rb.font.size = Pt(10.5)
    shade(cells[0], 'F2F2F2')
    for c in cells: c.paragraphs[0].paragraph_format.line_spacing = 1.1
rowt('Inventors', '1. Soumya Das\n2. Vipul Virsinghbhai Gamit\n3. Jagyanseni Paikaraya\n4. Tanisha\n5. Bharat Tank')
rowt('Inventor affiliations', 'Inventors 1–4: Department of Computer Applications (MCA), Faculty of IT and Computer Science, Parul University, Vadodara, Gujarat, India.\nInventor 5: Department of Computer Science and Engineering, PIET, Vadodara, Gujarat, India.')
rowt('Correspondence', 'vipulbhai.gamit26053@paruluniversity.ac.in')
rowt('Applicant / assignee', 'To be confirmed by the inventors and their institutions (ownership and assignment have not been determined in this draft).')
rowt('Field of invention', 'Computer-implemented medical image analysis; explainable artificial intelligence; dermatology decision support.')
rowt('Proposed classification (indicative only)', 'G06T 7/00; G06N 3/08; G06V 10/82; G16H 30/40; A61B 5/00 (to be verified by the agent).')
rowt('Disclosure warning', 'The underlying study has been prepared for conference submission (ICCSDI 2026). Any publication, preprint or presentation before filing may prejudice novelty; grace-period rules differ by jurisdiction and are narrow. Confirm the disclosure history with a patent agent before any further disclosure, and file first.')
rowt('Document status', 'Draft specification with claims, abstract and drawings. Prior-art search, novelty/inventive-step assessment and jurisdiction-specific formalities have not been carried out.')
doc.add_page_break()

# ============================================================================== TITLE + FIELD
H('TITLE OF THE INVENTION', 12, True, 0)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run('SYSTEM AND METHOD FOR NULL-CONTROLLED RELIABILITY ASSESSMENT OF SALIENCY EXPLANATIONS IN DERMOSCOPIC SKIN-LESION CLASSIFICATION'); r.bold = True
H('1. FIELD OF THE INVENTION')
P("The present invention relates to computer-implemented analysis of medical images, and more particularly to a system and method that assess whether a visual explanation (saliency map) produced for a neural-network classification of a dermoscopic skin-lesion image is spatially reliable, functionally faithful and reproducible, by scoring the explanation against a lesion segmentation mask while correcting for the geometric dependence of mask-overlap scores on lesion size, and that combines the resulting indicators with a calibrated confidence and a domain-shift indicator to control how the explanation is presented to a user.")
H('2. BACKGROUND OF THE INVENTION')
P("Deep neural networks classify dermoscopic images of pigmented skin lesions with high aggregate accuracy. To support clinical use, classifiers are commonly accompanied by a saliency map, for example a class-activation map computed by gradient-weighted methods such as Grad-CAM or Grad-CAM++, that is overlaid on the image to indicate the regions that influenced the prediction. Such maps are, however, produced without any indication of whether they are trustworthy for the particular image being viewed: a map may fall largely outside the lesion, may not reflect the evidence actually used by the model, or may change from one stochastic evaluation to the next.")
P("A known way to quantify spatial agreement is to compute the fraction of the total saliency mass that falls inside a reference region, such as an expert lesion segmentation mask (sometimes termed relevance-mass accuracy). The inventors have identified, and demonstrated experimentally on 990 held-out test images and five independently trained models, a previously unaddressed defect of this score when applied to lesions of different sizes. Because the score equals the area fraction of the mask when the saliency is uniform, and has the same expected value when the saliency is placed independently of the mask, the score increases steeply with lesion size by construction. In the experiments described below, the correlation between the raw score and the lesion-area fraction was between " + f"{rho_obs.min():.2f} and {rho_obs.max():.2f}" + " across five models, yet a correlation of " + f"{rho_null.min():.2f} to {rho_null.max():.2f}" + " was obtained when the same maps were paired with masks of other, unrelated images. A user or downstream rule that reads the raw score would therefore conclude that small lesions are explained poorly and large lesions well, largely irrespective of the model's actual behaviour, and a threshold on the raw score would systematically suppress explanations for small lesions.")
P("Further, the attainable maximum of the score is itself limited for small lesions when the saliency map is derived from a coarse convolutional feature map: for a 7 × 7 target layer, a map that is perfectly aligned with the lesion at that resolution reaches a median score of only about " + f"{sm['oracle']['med_small']:.2f}" + " for small lesions (area below 5% of the image) in the experiments described below.")
P("Spatial agreement is moreover not the same property as faithfulness (whether the prediction depends on the regions marked as important) or as reproducibility. In the experiments, the raw score was negatively or inconsistently related to insertion and deletion faithfulness measures, whereas an area-corrected score was moderately related to insertion faithfulness. Existing explanation-evaluation approaches (for example region-perturbation curves, deletion and insertion metrics, sanity checks and retraining-based benchmarks) address individual aspects, but the inventors are not aware of a system that (i) corrects a mask-overlap score for lesion geometry using image-matched reference and ceiling baselines, (ii) reports it together with faithfulness, stability, calibrated confidence and a domain-shift indicator, and (iii) uses the combination to control the presentation of the explanation. Such a system is needed to prevent explanations from conveying unwarranted assurance and to avoid size-biased suppression of explanations.")
H('3. OBJECTS OF THE INVENTION')
P("It is an object of the invention to provide a computer-implemented method and system that produce, for a saliency map generated for a dermoscopic image, a localization reliability indicator that is corrected for the dependence of mask-overlap scores on lesion area and on the spatial resolution of the saliency map.")
P("It is a further object to provide a combined trust report in which localization, faithfulness, stability, calibrated confidence and a domain-shift indicator are computed as separate quantities and are not conflated.")
P("It is a further object to provide an explanation-presentation controller that displays, annotates, suppresses or defers the explanation on the basis of the combined trust report without size-biased suppression.")
P("It is a further object to provide a method of decomposing the performance loss of a classifier evaluated on an external image collection into a population component, a label-space component, a class-prevalence component and a residual component.")
H('4. SUMMARY OF THE INVENTION')
P("According to a first aspect, a computer-implemented method comprises: receiving a dermoscopic image and a binary lesion mask; obtaining a predicted class produced by a neural-network classifier for the image; generating a saliency map for the predicted class from a target layer of the classifier, the target layer having a spatial resolution; computing a mask-overlap score as the ratio of saliency mass inside the lesion mask to total saliency mass; computing a reference expectation of the mask-overlap score that the lesion mask would obtain with saliency maps unrelated to the image; computing an excess localization value as the difference between the mask-overlap score and the reference expectation; and outputting a localization reliability indicator based at least on the excess localization value.")
P("In embodiments, the reference expectation is the mean mask-overlap obtained by pairing the lesion mask with saliency maps generated for other images of a reference set (a cross-image shuffled null), which preserves the classifier's spatial prior, and/or the mean obtained with spatially shifted copies of the saliency map of the same image (a shift null). In embodiments, a ceiling value is computed as the mask-overlap obtained by an idealised map derived from the lesion mask itself by pooling to the target-layer resolution and upsampling, and the indicator includes the ratio of the score to the ceiling. In embodiments, a permutation procedure re-pairs saliency maps and masks over a set of images to determine whether an observed association between mask-overlap scores and lesion area exceeds the association obtained under random pairing.")
P("In further embodiments, the method computes an insertion gain and a deletion gain that compare patch-wise insertion and deletion of image regions in order of saliency against random order, using the calibrated probability of the predicted class; computes a stability value as the mean pairwise correlation of saliency maps obtained under stochastic forward passes; applies temperature scaling to the classifier's outputs; and computes a domain-shift indicator. A trust aggregator combines these quantities, and an output interface displays, annotates, suppresses or defers the explanation.")
P("According to further aspects there are provided a system comprising at least one processor and memory storing instructions that, when executed, perform the method, and a non-transitory computer-readable medium storing such instructions. According to a further aspect there is provided a method of decomposing an external-evaluation performance drop.")
H('5. BRIEF DESCRIPTION OF THE DRAWINGS')
for f, d in [('FIG. 1', 'is a block diagram of a system according to an embodiment of the invention.'),
             ('FIG. 2', 'is a flowchart of a method for computing a localization reliability indicator and a trust report according to an embodiment.'),
             ('FIG. 3', 'shows experimental results of the lesion-size audit: observed mask-overlap score against lesion area together with the shuffled-map reference, the shift reference, the ceiling and the uniform expectation; the excess localization value against lesion area; and the size-stratum comparison.'),
             ('FIG. 4', 'shows experimental deletion and insertion curves for saliency order, random order and mask-derived orders, and the relation of excess localization to the insertion and deletion gains.'),
             ('FIG. 5', 'shows the distribution of stochastic saliency stability values and their relation to excess localization.'),
             ('FIG. 6', 'shows a stepwise decomposition of the accuracy drop on an external image collection and an adaptation ladder.'),
             ('FIG. 7', 'shows example outputs: dermoscopic images, saliency overlays with the lesion outline, raw scores, reference values and excess values.'),
             ('FIG. 8', 'illustrates the separate trust dimensions computed by the system.')]:
    P(f"{f} {d}")
para_no_start_detail = para_no[0]
H('6. DETAILED DESCRIPTION OF THE INVENTION', page_break=True)
P("The following description sets out embodiments by way of example. Specific values (layer, patch size, number of passes and the like) are those of the experimentally validated embodiment and are not limiting; the claims define the scope of protection. Reference numerals refer to FIG. 1 and FIG. 2.")
H('6.1 System overview', 11, before=6)
P("Referring to FIG. 1, a system 10 comprises an image source 100 (for example a dermoscope, a picture-archiving system or a file store), a pre-processing unit 102, a neural-network classifier 104, a confidence calibration unit 106, a saliency-map generator 108, a lesion-mask source 110, a localization scorer 112, a reference generator 114, a faithfulness evaluator 116, a stability evaluator 118, an optional domain-shift monitor 120, a trust aggregator 122 and an output interface 124. The units may be implemented as software modules executing on one or more processors (for example a central processing unit and a graphics processing unit) with associated memory, or distributed over a network, and may be integrated in a dermoscopy workstation, a tele-dermatology platform or a clinical decision-support application.")
H('6.2 Pre-processing and classification', 11, before=6)
P("The pre-processing unit 102 removes hair artefacts from the dermoscopic image, in one embodiment by a black-hat morphological filter with a 9 × 9 rectangular kernel and a threshold of 10, a morphological opening with a 5 × 5 elliptical kernel to separate thin structures from blob-like pigment structures, removal of connected components of fewer than 8 pixels, a 3 × 3 dilation, and inpainting (Telea method, radius 3); resizes the image to 224 × 224 pixels; and normalises it with ImageNet channel statistics. The classifier 104 is, in the validated embodiment, a DenseNet121 pretrained on ImageNet whose classifier is replaced by dropout (p = 0.3), a 512-unit linear layer, batch normalisation, a rectified-linear unit, dropout (p = 0.2) and a seven-way linear layer for the classes actinic keratosis/intraepithelial carcinoma, basal cell carcinoma, benign keratosis-like lesion, dermatofibroma, melanoma, melanocytic nevus and vascular lesion. In the validated embodiment the classifier is trained with a class-weighted focal loss (γ = 2) plus 0.1 times a label-smoothed cross-entropy, a class-balanced sampler, AdamW with discriminative learning rates (1 × 10⁻⁵ to 3 × 10⁻⁴), and a complete 60-epoch one-cycle schedule in which the learning rate is annealed to approximately zero, the checkpoint being the epoch with the highest macro-F1 on a selection subset of a validation set that is disjoint from a calibration subset.")
P("The confidence calibration unit 106 divides the logits by a scalar temperature T fitted on the calibration subset by minimising negative log-likelihood, giving calibrated class probabilities p(k) = exp(z_k/T) / Σ_j exp(z_j/T). The temperature does not change the predicted class.")
H('6.3 Saliency map and mask-overlap score', 11, before=6)
P("The saliency-map generator 108 computes, for the predicted class c, a class-activation map H from a target layer of the classifier 104 — in the validated embodiment the last convolution of the fourth dense block (7 × 7 spatial resolution for a 224 × 224 input) using the Grad-CAM++ weighting — upsamples it bilinearly to the image size and normalises it to the range [0, 1]. The lesion-mask source 110 supplies a binary lesion mask M for the image, for example an expert or semi-automatic segmentation, or the output of a segmentation network. The localization scorer 112 computes the mask-overlap score (Equation 1):")
P("LSAS(x) = Σ_{p∈M} H_p / Σ_{p∈Ω} H_p        (1)", numbered=False).alignment = WD_ALIGN_PARAGRAPH.CENTER
P("where Ω is the set of all pixels. With a = |M|/|Ω| the lesion-area fraction, LSAS = a for spatially uniform H, and E[LSAS] = a for any H whose location is independent of M and uniformly distributed over the image, regardless of how concentrated H is. LSAS is therefore expected to increase with a even for maps that carry no information about the image.")
H('6.4 Reference generator: null and ceiling', 11, before=6)
P("The reference generator 114 provides one or more of the following reference quantities. (a) **Cross-image shuffled null** E_shuf(M_i): the mean of the mask-overlap scores obtained by pairing the mask M_i with the saliency maps H_j of every other image j of a reference set. This null preserves the spatial prior of the classifier and saliency method (including the tendency of maps to be central in dermoscopic images) while destroying the image-specific link. (b) **Shift null** E_shift(H_i, M_i): the mean over n (for example 100) random circular shifts of H_i against M_i, together with its standard deviation and a one-sided permutation p-value (rank of the observed score among the shifted scores). (c) **Ceiling** C_max(M_i): the mask-overlap score of an idealised map obtained by average-pooling M_i to the target-layer resolution (for example 7 × 7), bilinearly upsampling it to the image size and normalising it to [0, 1], that is, the score a perfectly lesion-aligned map with the resolution of the target layer would obtain. (d) The uniform reference a.")
P("The localization scorer 112 outputs the excess localization value (Equation 2) and, optionally, the attainable fraction (Equation 3), a z-score against the shift null and the rank of the observed score among the shuffled scores:")
P("e(x) = LSAS(x) − E_shuf(M)        (2)", numbered=False).alignment = WD_ALIGN_PARAGRAPH.CENTER
P("f(x) = LSAS(x) / C_max(M)        (3)", numbered=False).alignment = WD_ALIGN_PARAGRAPH.CENTER
P("In one embodiment the scorer additionally computes an area-independent statistic, namely the receiver-operating-characteristic area under the curve of the saliency values as a predictor of mask membership, whose value under chance is 0.5 irrespective of lesion area.")
H('6.5 Association test against random pairing', 11, before=6)
P("For a set of images, the system may determine whether the observed association between mask-overlap score and lesion area exceeds the association explained by geometry. For B re-pairings (for example B = 2,000) of saliency maps with lesion masks, in which no map is paired with its own mask, the Spearman correlation between the re-paired scores and the lesion areas (and, optionally, the Kruskal–Wallis statistic across size strata such as area < 5%, 5–15% and ≥ 15%) is computed to form a reference distribution; the observed statistic is compared with this distribution. An observed statistic that does not exceed the reference distribution indicates that the size dependence of the score is geometric and should not be interpreted as a property of the classifier. The output of this test may be used to qualify the reliability indicator or to disable size-based thresholds.")
H('6.6 Faithfulness evaluator', 11, before=6)
P("The faithfulness evaluator 116 partitions the (normalised) input into non-overlapping patches (in the validated embodiment 14 × 14 = 196 patches of 16 × 16 pixels) and scores each patch by the mean of H over the patch. In a deletion pass, patches are replaced by a baseline (the mean image, i.e. zero after normalisation, or a blurred copy of the image) in order of decreasing saliency; in an insertion pass, the baseline image is progressively replaced by the original patches in the same order. After each of a number of steps (in the validated embodiment 0, 14, 28, …, 196 patches, i.e. 15 points) the calibrated probability of the predicted class is recorded, and the area under each curve is computed. The same procedure is applied to random patch orders (for example the mean of three random permutations) and, optionally, to orders derived from the lesion mask (lesion-covered patches first, and background patches first). The deletion gain is DG = AUC_del(random) − AUC_del(saliency) and the insertion gain is IG = AUC_ins(saliency) − AUC_ins(random); positive values indicate that the saliency order is more informative than chance. The evaluator does not equate a high localization value with faithfulness: the two quantities are reported separately and their relation may be reported as a correlation over a set of images.")
H('6.7 Stability evaluator', 11, before=6)
P("The stability evaluator 118 obtains N (for example 15) saliency maps for the predicted class of the deterministic model with dropout layers active (independent masks per pass, batch-normalisation in evaluation mode) and computes the stability value as the mean pairwise Pearson correlation of the flattened maps. A stability value below a threshold (for example 0.5) marks the explanation as unstable. The dropout layers must be re-enabled after any explanation library constructor that resets the model to evaluation mode; otherwise all passes are identical and the stability value is trivially 1.")
H('6.8 Domain-shift monitor', 11, before=6)
P("The optional domain-shift monitor 120 compares features of an incoming image, or of a batch of incoming images, with features of a source reference set — for example by a domain classifier (a linear classifier separating source from incoming features, whose cross-validated AUROC near 0.5 indicates similarity and near 1.0 indicates a strong shift) or by a Fréchet distance between feature distributions — and raises a domain flag when the measure exceeds a threshold. The flag is passed to the trust aggregator 122 so that explanations and confidences obtained outside the validated imaging domain are marked accordingly.")
H('6.9 Trust aggregator and output interface', 11, before=6)
P("The trust aggregator 122 outputs a trust report containing, as separate fields, the predicted class and calibrated confidence, the mask-overlap score, the excess localization value and attainable fraction, the insertion and deletion gains, the stability value and the domain flag, and derives a reliability indicator (for example categorical: reliable, uncertain, unreliable) by predetermined rules over these fields. The output interface 124 displays, annotates, suppresses or defers the explanation according to the indicator. In particular, the decision to suppress or flag a saliency overlay is based on the excess localization value and the other fields and not on the raw mask-overlap score or a threshold applied to it, so that small lesions are not systematically flagged for geometric reasons. Where a lesion mask is unavailable, the mask may be produced by a segmentation network within the lesion-mask source 110, and the reliability indicator is then qualified as depending on the segmentation.")
H('6.10 Method', 11, before=6)
P("Referring to FIG. 2, in step S200 the dermoscopic image and lesion mask are received, and the predicted class and calibrated confidence obtained. In S202 the saliency map is generated. In S204 the mask-overlap score is computed. In S206 the reference expectation is computed and in S208 the ceiling. In S210 the excess localization value and attainable fraction are computed. In S212 the insertion and deletion gains and the stability value are computed. In S214 the values are aggregated into the reliability indicator, and in S216 the explanation is displayed, suppressed, annotated or deferred.")
H('6.11 Decomposition of an external-evaluation performance drop', 11, before=6)
P("In a further aspect, for a classifier trained on a source collection and evaluated on an external collection whose classes map to a subset of the source classes, the performance drop is decomposed by computing, from the same model outputs: (A) accuracy on the source test set; (B) accuracy on source test images restricted to the mappable classes with the full output space (evaluation-population effect); (C) accuracy with the output restricted to the mappable classes (label-space effect); (P) the expected accuracy obtained by re-weighting the per-class recalls of (C) to the class prevalence of the external collection (prevalence effect); and (E) accuracy on the external collection; the difference between (P) and (E) is the residual attributable to appearance and label-semantics differences. Macro-averaged metrics are computed over the classes present in the evaluated labels, since counting absent classes as zero inflates the apparent drop. The residual may be further characterised by (i) a domain classifier and Fréchet distance between feature distributions and (ii) an adaptation ladder in which, on a patient-level split of the external collection, a class-bias correction and then a linear classifier on frozen features are fitted on a minority of patients and evaluated on the remainder.")
H('6.12 Experimental embodiments and validation', 11, before=6)
P(f"The system was implemented and evaluated on the public HAM10000 collection (10,015 dermoscopic images of 7,470 lesions; lesion-level split 8,020 / 1,005 / 990 images for training, validation and test) with the expert lesion segmentations supplied with that collection, and on the public PAD-UFES-20 smartphone collection (2,298 images; 2,106 after mapping of five diagnostic classes). The entire training and evaluation pipeline was repeated for five independently seeded runs (seeds 42–46). Values below are mean ± standard deviation over the five runs unless stated. These are computational results on retrospective public data and do not establish clinical performance.")
T(['Quantity', 'Result (5 runs)'],
  [['Test accuracy / balanced accuracy / macro-F1', f"{pm(agg['acc']['mean'], agg['acc']['sd'])} / {pm(agg['bacc']['mean'], agg['bacc']['sd'])} / {pm(agg['f1']['mean'], agg['f1']['sd'])}"],
   ['Macro-AUROC / macro-PR-AUC / MCC', f"{pm(agg['auroc']['mean'], agg['auroc']['sd'])} / {pm(agg['prauc']['mean'], agg['prauc']['sd'])} / {pm(agg['mcc']['mean'], agg['mcc']['sd'])}"],
   ['Expected calibration error, after / before temperature scaling', f"{pm(agg['ece']['mean'], agg['ece']['sd'], 3)} / {pm(agg['ece_uncal']['mean'], agg['ece_uncal']['sd'], 3)}"],
   ['Raw mask-overlap score vs. lesion area (Spearman ρ)', f"{rho_obs.min():.2f} to {rho_obs.max():.2f} across runs"],
   ['Same, mask paired with maps of other images (null), mean ρ', f"{rho_null.min():.2f} to {rho_null.max():.2f} across runs"],
   ['Observed ρ minus null ρ', f"{gap.min():+.3f} to {gap.max():+.3f} (observed did not exceed the null in {int((ac('perm_null', 'p_greater') >= 0.05).sum())} of {K} runs)"],
   ['Median excess localization value e (all 990 test images)', f"{e_med.min():+.3f} to {e_med.max():+.3f}"],
   ['Spearman ρ of excess value e with lesion area', f"{e_area.min():+.2f} to {e_area.max():+.2f} (negative in every run)"],
   ['Ceiling C_max, median for small / medium / large lesions (reference run)', f"{sm['oracle']['med_small']:.2f} / {sm['oracle']['med_medium']:.2f} / {sm['oracle']['med_large']:.2f}"],
   ['Insertion gain IG (median; images with IG > 0)', f"{ig.min():.3f} to {ig.max():.3f}; {ig_f.min():.0f}–{ig_f.max():.0f}%"],
   ['Deletion gain DG (median)', f"{dg.min():+.3f} to {dg.max():+.3f} (negative in {int((dg < 0).sum())} of {K} runs)"],
   ['Spearman ρ: raw score vs. IG / raw score vs. DG', f"{l_ig.min():+.2f} to {l_ig.max():+.2f} / {l_dg.min():+.2f} to {l_dg.max():+.2f}"],
   ['Spearman ρ: excess value e vs. IG', f"{e_ig.min():+.2f} to {e_ig.max():+.2f}"],
   ['Stability value, median (300 images, 15 passes)', f"{sts_m.min():.2f} to {sts_m.max():.2f}; ρ with excess value {sts_e.min():+.2f} to {sts_e.max():+.2f}; ρ with raw score {sts_l.min():+.2f} to {sts_l.max():+.2f}"]],
  'TABLE 1. Experimental results of the validated embodiment.')
P(f"__Lesion-size audit (FIG. 3).__ The raw mask-overlap score increased steeply with lesion area (medians {sm['lsas']['med_small']:.2f}, {sm['lsas']['med_medium']:.2f} and {sm['lsas']['med_large']:.2f} for small, medium and large lesions in the reference run), but the cross-image shuffled reference reproduced this gradient almost exactly (medians {sm['null_shuf_mean']['med_small']:.2f}, {sm['null_shuf_mean']['med_medium']:.2f} and {sm['null_shuf_mean']['med_large']:.2f}). A model-independent simulation on the real masks gave Spearman correlations with area of {sim['uniform']['rho_area']:.2f} for uniform saliency, {sim['diffuse_random_location(sigma=60)']['rho_area']:.2f} for diffuse saliency at a random location and {sim['concentrated_random_location(sigma=12)']['rho_area']:.2f} for concentrated saliency at a random location. After subtraction of the shuffled reference the excess value was small and decreased with lesion size in every run. The invention thus prevents the misreading that small lesions are explained worse.")
P(f"__Faithfulness (FIG. 4).__ Insertion of saliency-ranked patches raised the calibrated probability of the predicted class substantially more than random insertion in every run, whereas the deletion evidence was inconsistent between runs. The raw score did not predict faithfulness, whereas the excess value was moderately and consistently related to the insertion gain. Reporting localization and faithfulness as separate fields, rather than one composite score, therefore avoids conveying unwarranted assurance.")
P(f"__Stability (FIG. 5).__ Median stability was {sts_m.min():.2f} to {sts_m.max():.2f}, and the stability value was moderately to strongly related to the excess value but only weakly and negatively to the raw score, which shows that the raw score conflates geometry with reproducibility.")
P(f"__External collection (FIG. 6).__ Source accuracy of {A7.mean():.3f} became {C5.mean():.3f} when restricted to the five mappable classes (population and label-space effects together {C5.mean() - A7.mean():+.3f}), an expected {Pv.mean():.3f} after re-weighting to the external class prevalence, and {E5.mean():.3f} observed on the external collection, so that the residual attributable to appearance and label semantics was {E5.mean() - Pv.mean():+.3f}. A domain classifier separated the collections with an AUROC of {dcauc.mean():.4f}. In the adaptation ladder (balanced accuracy on evaluation patients), zero-shot, class-bias-corrected and linear-probe values were {lad('zero_shot', 'bacc').mean():.3f}, {lad('bias_only', 'bacc').mean():.3f} and {lad('probe_1024', 'bacc').mean():.3f}, respectively.")
P("Example outputs of the system for small and large lesions, for correct and incorrect predictions, are shown in FIG. 7, and the separate trust dimensions computed by the system are illustrated in FIG. 8.")
H('6.13 Advantages and technical effect', 11, before=6)
P("The technical effect of the invention is that the reliability indicator of an explanation produced by a neural-network image classifier is made independent of the geometry of the reference region and of the resolution of the explanation layer, so that explanation display control does not depend on lesion size for reasons unrelated to the classifier. The invention further provides, from the same model outputs, separately computed faithfulness, stability, calibration and domain-shift measures, which reduces the risk that a user infers functional support for a prediction from spatial overlap alone.")
H('6.14 Variations', 11, before=6)
P("The invention is not limited to dermoscopy or to DenseNet121 and Grad-CAM++. The classifier may be any convolutional or transformer classifier; the saliency generator may be any method producing a spatial relevance map at a known resolution (for example Grad-CAM, Score-CAM, attention rollout or layer-wise relevance propagation); the reference region may be any segmentation of a lesion or organ; the null may be built from any reference set of images with masks; the patch size, number of steps, baselines, number of stochastic passes, number of shifts and thresholds may vary; and the system may be applied to other imaging modalities in which mask-overlap scores of explanations are used. The excess value may be expressed as a difference, ratio or standardised score, and may be combined with the attainable fraction by rules, a logistic model or another classifier. Such variations are within the scope of the claims.")
H('7. LIMITATIONS OF THE DISCLOSED VALIDATION', 12)
P("The experiments used two public retrospective collections, one architecture and one saliency method at one layer, five training seeds and a single data split; no reader study, prospective study or regulatory evaluation was performed. Deletion-based faithfulness was not consistent across training runs and depended on the baseline, and the skin-tone-proxy analyses of the underlying study were inconclusive because of the small number of darker-proxy images. Nothing in this specification asserts clinical safety, diagnostic performance in practice or fairness across skin tones.")

# ============================================================================== CLAIMS
H('WE CLAIM:', 13, True, page_break=True)
claims = [
 "1. A computer-implemented method for assessing the reliability of a saliency explanation of a neural-network classification of a dermoscopic image, the method comprising: (a) receiving the dermoscopic image and a binary lesion mask of the image; (b) obtaining, from a neural-network classifier, a predicted class for the image; (c) generating a saliency map for the predicted class from a target layer of the classifier, the target layer having a spatial resolution; (d) computing a mask-overlap score as a ratio of the saliency mass located inside the lesion mask to the total saliency mass; (e) computing a reference expectation of the mask-overlap score for the lesion mask under saliency maps that are unrelated to the dermoscopic image; (f) computing an excess localization value as a difference between the mask-overlap score and the reference expectation; and (g) outputting a localization reliability indicator based at least on the excess localization value.",
 "2. The method of claim 1, wherein computing the reference expectation comprises pairing the lesion mask with saliency maps generated by the classifier for other images of a reference set and averaging the resulting mask-overlap scores.",
 "3. The method of claim 1, wherein computing the reference expectation comprises computing mask-overlap scores between the lesion mask and a plurality of spatially shifted copies of the saliency map, and the method further comprises computing at least one of a standardised score and a permutation p-value of the mask-overlap score against said scores.",
 "4. The method of claim 1, further comprising computing a ceiling value as a mask-overlap score of a map derived from the lesion mask by pooling the lesion mask to the spatial resolution of the target layer and upsampling the pooled mask to the size of the image, and outputting an attainable fraction equal to the ratio of the mask-overlap score to the ceiling value.",
 "5. The method of claim 1, further comprising computing an area-independent localization statistic being an area under a receiver-operating-characteristic curve of saliency values as a predictor of membership of the lesion mask.",
 "6. The method of claim 1, further comprising, for a set of images: computing a statistic of association between mask-overlap scores and lesion-area fractions; repeating the computation for a plurality of re-pairings of saliency maps with lesion masks of different images to obtain a reference distribution of the statistic; and determining that a dependence of the mask-overlap score on lesion size is geometric when the statistic does not exceed the reference distribution.",
 "7. The method of claim 1, further comprising: dividing the image into non-overlapping patches; ranking the patches by saliency; computing a first curve of the calibrated probability of the predicted class as patches are progressively inserted into a baseline image in order of the ranking and a second curve as patches are progressively deleted from the image in order of the ranking; computing corresponding curves for random patch orders; and outputting an insertion gain and a deletion gain as differences between areas under the curves for the ranked order and the random order.",
 "8. The method of claim 7, wherein the baseline image is a mean image or a blurred copy of the image, and the method further comprises computing insertion and deletion curves for orders derived from the lesion mask in which patches covered by the lesion are ordered before or after patches covered by background.",
 "9. The method of claim 7, further comprising reporting a correlation, over a plurality of images, between the excess localization value and at least one of the insertion gain and the deletion gain, and reporting the localization reliability indicator and the insertion and deletion gains as separate fields.",
 "10. The method of claim 1, further comprising: obtaining a plurality of saliency maps for the predicted class from the classifier with dropout layers active and batch normalisation in evaluation mode; computing a stability value as a mean pairwise correlation of the plurality of saliency maps; and marking the saliency explanation as unstable when the stability value is below a threshold.",
 "11. The method of claim 1, further comprising calibrating outputs of the classifier by dividing logits by a scalar temperature fitted on a calibration set that is disjoint from a set used for selecting a trained checkpoint of the classifier, and outputting a calibrated confidence of the predicted class.",
 "12. The method of claim 1, further comprising computing a domain-shift indicator by comparing features of the image, or of a batch of images including the image, with features of a source reference set, and outputting the domain-shift indicator together with the localization reliability indicator.",
 "13. The method of claim 1, further comprising aggregating the localization reliability indicator, an insertion or deletion gain, a stability value, a calibrated confidence and a domain-shift indicator into a trust report in which each is a separate field, and displaying, annotating, suppressing or deferring the saliency explanation on a display device according to the trust report, wherein the suppressing or deferring is not determined by a threshold applied to the mask-overlap score alone.",
 "14. The method of claim 1, wherein the lesion mask is generated by a segmentation neural network applied to the dermoscopic image, and the localization reliability indicator is qualified as dependent on the segmentation.",
 "15. The method of claim 1, wherein the classifier is a convolutional neural network, the saliency map is a gradient-weighted class-activation map, the target layer is a last convolutional layer of a dense block, and the reference expectation is computed for a reference set of at least 100 images.",
 "16. A system for assessing the reliability of a saliency explanation of a neural-network classification of a dermoscopic image, comprising: at least one processor; and memory storing instructions that, when executed by the at least one processor, cause the system to perform the method of any one of claims 1 to 15.",
 "17. The system of claim 16, further comprising a display device and an output interface configured to display, annotate, suppress or defer the saliency explanation according to the localization reliability indicator.",
 "18. A non-transitory computer-readable storage medium storing instructions that, when executed by a processor, cause the processor to perform the method of any one of claims 1 to 15.",
 "19. A computer-implemented method of characterising a performance drop of a neural-network image classifier trained on a source collection when evaluated on an external collection whose classes correspond to a subset of the classes of the source collection, the method comprising: computing, from outputs of the classifier, (A) a source accuracy on a source test set; (B) an accuracy on source test images restricted to images of the corresponding classes; (C) an accuracy on said restricted images with the output space restricted to the corresponding classes; (P) an expected accuracy obtained by weighting per-class recalls of (C) by class prevalence of the external collection; and (E) an accuracy on the external collection; and outputting differences between successive ones of (A), (B), (C), (P) and (E) as an evaluation-population component, a label-space component, a class-prevalence component and a residual component, respectively.",
 "20. The method of claim 19, wherein macro-averaged metrics are computed over classes present in the labels of the evaluated collection, and the method further comprises evaluating, on a patient-level split of the external collection, recovery of performance by a class-bias correction and by a linear classifier fitted on frozen features of the classifier, and computing a domain-classifier area under the curve between source and external features.",
]
for c in claims:
    p = doc.add_paragraph(c); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after = Pt(9)

# ============================================================================== ABSTRACT
H('ABSTRACT', 13, True, page_break=True)
abstract = ("A system and method assess the reliability of a saliency explanation for a neural-network classification of a dermoscopic image. A mask-overlap score, the fraction of saliency mass inside a lesion mask, is computed and corrected for its geometric dependence on lesion size by subtracting a reference expectation obtained with saliency maps of other images or shifted maps, yielding an excess localization value; a ceiling based on the mask pooled to the resolution of the saliency layer gives an attainable fraction. Faithfulness (insertion and deletion gains against random order), stability under stochastic passes, calibrated confidence and a domain-shift indicator are computed as separate fields of a trust report that controls whether the explanation is displayed, annotated, suppressed or deferred. A method of decomposing an external-evaluation performance drop into population, label-space, prevalence and residual components is also disclosed.")
p = doc.add_paragraph(abstract); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
c = doc.add_paragraph(); r = c.add_run(f'(Word count: {len(abstract.split())})'); r.italic = True; r.font.size = Pt(10)
H('Representative figure for the abstract: FIG. 1', 11)

# ============================================================================== DRAWINGS
sec2 = doc.add_page_break()
H('DRAWINGS', 13, True, before=0)
p = doc.add_paragraph('The drawings follow on separate sheets. FIG. 1–2 are original line drawings; FIG. 3–8 are grayscale reproductions of experimental results produced by the disclosed system.'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
FIG(PD / 'patent_fig1_system.png', 'FIG. 1 — System 10', 15.5)
FIG(PD / 'patent_fig2_flow.png', 'FIG. 2 — Method flowchart', 11.5)
FIG(PD / 'patent_fig3.png', 'FIG. 3 — Lesion-size audit: observed score, references, ceiling and excess value', 16)
FIG(PD / 'patent_fig4.png', 'FIG. 4 — Deletion/insertion curves and relation of excess localization to faithfulness gains', 16)
FIG(PD / 'patent_fig5.png', 'FIG. 5 — Stochastic saliency stability', 13)
FIG(PD / 'patent_fig6.png', 'FIG. 6 — Decomposition of the external-collection accuracy drop and adaptation ladder', 16)
FIG(PD / 'patent_fig7.png', 'FIG. 7 — Example outputs for small and large lesions', 16)
FIG(PD / 'patent_fig8.png', 'FIG. 8 — Separate trust dimensions computed by the system', 13)
doc.save(OUT / 'Patent_Specification_TrustXAI_Derm.docx')
print('patent built; paragraphs:', para_no[0], 'claims:', len(claims), 'abstract words:', len(abstract.split()))
